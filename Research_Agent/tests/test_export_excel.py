"""Tests for export_excel.py -- the dashboard's on-demand "Excel: Companies
& Contacts" download. Stub-based, same style as test_export.py: no network,
db.* is monkeypatched directly so these lock in the read path (get_scores /
contacts_by_domain / companies_by_domain) and the workbook's shape rather
than any live Supabase behavior.

Critical constraint under test: this must stay a LEAN list -- one row per
company (or per contact, for companies with 2+ contacts), list-level fields
only, never the full research report (no pain_summary/explanation/reasons
columns -- see test_excel_has_no_full_report_fields).
"""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import load_workbook

from csg import db, export_excel
from csg.models import Company

SCORE_ROW = {
    "domain": "acme.com", "run_id": "t1", "score": 71.5, "tier": "A",
    "dimensions": {"Hiring Need": 30.0, "Fill Pain": 15.0},
    "explanation": ["+30 Hiring Need: 12 relevant open positions"],
    "scored_at": datetime(2026, 9, 16).isoformat(),
    "metrics": {"req_count": 12},
    "name": "Acme Data", "region": "GA", "industry": "Software / AI", "hq_state": "GA",
    "employee_count": {"value": 400, "source": "gemini", "method": "estimated", "confidence": "medium"},
    "revenue": {"value": "$25M", "source": "gemini", "method": "estimated", "confidence": "medium"},
    "funding": {"value": "$40M Series B", "source": "gemini", "method": "reported", "confidence": "high"},
    "ats_provider": "greenhouse", "ats_token": "acme",
    "hiring_signal": "confirmed", "job_discovery_method": "ats",
}

ACME_COMPANY = Company(domain="acme.com", name="Acme Data", careers_url="https://acme.com/careers")


class _StubConn:
    def close(self): ...


@pytest.fixture
def stub(monkeypatch):
    monkeypatch.setattr(db, "get_scores", lambda c, r: [SCORE_ROW])
    monkeypatch.setattr(db, "latest_scores_per_domain", lambda c: [SCORE_ROW])
    monkeypatch.setattr(db, "contacts_by_domain", lambda c, d: {
        "acme.com": [{
            "email": "jane@acme.com", "first_name": "Jane", "last_name": "Doe",
            "position": "VP Talent", "linkedin_url": "https://linkedin.com/in/janedoe",
        }],
    })
    monkeypatch.setattr(db, "companies_by_domain", lambda c, d: {"acme.com": ACME_COMPANY})
    return _StubConn()


def _load(content: bytes):
    from io import BytesIO
    return load_workbook(BytesIO(content))


def test_headers_match_declared_columns(stub):
    wb = export_excel.build_workbook(stub, "t1")
    ws = wb.active
    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header_row == [h for _, h in export_excel.COLUMNS]


def test_one_row_per_company_with_single_contact(stub):
    wb = export_excel.build_workbook(stub, "t1")
    ws = wb.active
    assert ws.max_row == 2  # header + 1 data row


def test_row_values_correct(stub):
    wb = export_excel.build_workbook(stub, "t1")
    ws = wb.active
    keys = [k for k, _ in export_excel.COLUMNS]
    row = {k: v for k, v in zip(keys, [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))])}
    assert row["company_name"] == "Acme Data"
    assert row["domain"] == "acme.com"
    assert row["tier"] == "A"
    assert row["score"] == 71.5
    assert row["qualification_status"] == "Highly Qualified (Tier A)"
    assert row["industry"] == "Software / AI"
    assert row["employee_count"] == 400  # provenance unwrapped to bare value
    assert row["revenue"] == "$25M"
    assert row["funding"] == "$40M Series B"
    assert row["ats_provider"] == "greenhouse"
    assert row["open_reqs"] == 12
    assert row["hiring_signal"] == "confirmed"
    assert row["careers_url"] == "https://acme.com/careers"
    assert row["outreach_ready"] is True
    assert row["contact_name"] == "Jane Doe"
    assert row["contact_title"] == "VP Talent"
    assert row["contact_email"] == "jane@acme.com"
    assert row["contact_linkedin"] == "https://linkedin.com/in/janedoe"


def test_provenance_envelope_never_leaks_into_cell(stub):
    """The Excel list must never carry a raw provenance dict -- only the
    unwrapped scalar (see models.provenance_value)."""
    wb = export_excel.build_workbook(stub, "t1")
    ws = wb.active
    row = [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))]
    for v in row:
        assert not isinstance(v, dict)


def test_excel_has_no_full_report_fields(stub):
    """Critical product constraint: the Excel export is a lean list, never
    a duplicate of the full research report. No column should carry
    scoring explanation, pain summary, or per-dimension breakdown text."""
    keys = {k for k, _ in export_excel.COLUMNS}
    forbidden = {"explanation", "pain_summary", "recommended_angle", "dimensions", "reasons", "metrics"}
    assert keys.isdisjoint(forbidden)


def test_multiple_contacts_produce_multiple_rows(stub, monkeypatch):
    monkeypatch.setattr(db, "contacts_by_domain", lambda c, d: {
        "acme.com": [
            {"email": "jane@acme.com", "first_name": "Jane", "last_name": "Doe", "position": "VP Talent"},
            {"email": "bob@acme.com", "first_name": "Bob", "last_name": "Smith", "position": "HR Director"},
        ],
    })
    wb = export_excel.build_workbook(stub, "t1")
    ws = wb.active
    assert ws.max_row == 3  # header + 2 contact rows for the same company
    emails = [row[0].value for row in ws.iter_rows(min_row=2, max_row=3,
              min_col=export_excel.COLUMNS.index(("contact_email", "Contact Email")) + 1,
              max_col=export_excel.COLUMNS.index(("contact_email", "Contact Email")) + 1)]
    assert emails == ["jane@acme.com", "bob@acme.com"]


def test_company_with_no_contacts_still_gets_one_row(stub, monkeypatch):
    monkeypatch.setattr(db, "contacts_by_domain", lambda c, d: {})
    wb = export_excel.build_workbook(stub, "t1")
    ws = wb.active
    assert ws.max_row == 2
    keys = [k for k, _ in export_excel.COLUMNS]
    row = {k: v for k, v in zip(keys, [c.value for c in next(ws.iter_rows(min_row=2, max_row=2))])}
    assert row["contact_email"] is None or row["contact_email"] == ""
    assert row["outreach_ready"] is False


def test_missing_optional_fields_handled_gracefully(stub, monkeypatch):
    """A company missing industry/region/revenue/funding/careers_url must
    not crash the export -- it should just render blank."""
    thin_row = dict(SCORE_ROW)
    for k in ("industry", "region", "hq_state", "revenue", "funding", "ats_provider"):
        thin_row[k] = None
    thin_row["employee_count"] = None
    monkeypatch.setattr(db, "get_scores", lambda c, r: [thin_row])
    monkeypatch.setattr(db, "latest_scores_per_domain", lambda c: [thin_row])
    monkeypatch.setattr(db, "companies_by_domain", lambda c, d: {})
    wb = export_excel.build_workbook(stub, "t1")
    ws = wb.active
    assert ws.max_row == 2  # did not crash, produced exactly one row


def test_empty_run_produces_headers_only_workbook(monkeypatch):
    """A run with zero companies must still produce a valid, well-formed
    workbook -- headers only, no data rows, never a crash."""
    monkeypatch.setattr(db, "get_scores", lambda c, r: [])
    monkeypatch.setattr(db, "contacts_by_domain", lambda c, d: {})
    monkeypatch.setattr(db, "companies_by_domain", lambda c, d: {})
    wb = export_excel.build_workbook(_StubConn(), "empty-run")
    ws = wb.active
    assert ws.max_row == 1
    header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header_row == [h for _, h in export_excel.COLUMNS]


def test_to_bytes_roundtrips_through_openpyxl(stub):
    """Validates the actual bytes returned to the API route are a real,
    readable .xlsx file -- not just that build_workbook() didn't raise."""
    content = export_excel.to_bytes(stub, "t1")
    assert isinstance(content, bytes)
    assert len(content) > 0
    wb = _load(content)
    ws = wb.active
    assert ws["A1"].value == "Company"
    assert ws["A2"].value == "Acme Data"


def test_large_run_performance_and_correctness(monkeypatch):
    """A 100-company run must complete quickly and produce exactly the
    right number of rows -- no dropped or duplicated companies."""
    import time

    rows = []
    for i in range(100):
        row = dict(SCORE_ROW)
        row["domain"] = f"company{i}.com"
        row["name"] = f"Company {i}"
        rows.append(row)
    monkeypatch.setattr(db, "get_scores", lambda c, r: rows)
    monkeypatch.setattr(db, "latest_scores_per_domain", lambda c: rows)
    monkeypatch.setattr(db, "contacts_by_domain", lambda c, d: {
        f"company{i}.com": [{"email": f"contact{i}@company{i}.com", "first_name": "A", "last_name": "B"}]
        for i in range(100)
    })
    monkeypatch.setattr(db, "companies_by_domain", lambda c, d: {})

    start = time.monotonic()
    wb = export_excel.build_workbook(_StubConn(), "big-run")
    elapsed = time.monotonic() - start

    ws = wb.active
    assert ws.max_row == 101  # header + 100 companies, one contact each
    assert elapsed < 10.0

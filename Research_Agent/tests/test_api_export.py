"""Tests for GET /runs/{run_id}/export/excel and GET /runs/{run_id}/export/pdf
-- the dashboard's two on-demand download endpoints. Same TestClient +
monkeypatch style as test_api_usage.py: db.connect and the underlying db.*
read functions are stubbed so these lock in the HTTP contract (status,
media type, Content-Disposition, and that the exporter modules are actually
invoked with the run_id from the URL) without touching a real Supabase
project or the pipeline.
"""

from __future__ import annotations

import zipfile
from datetime import datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from csg import api, db
from csg.models import Company

SCORE_ROW = {
    "domain": "acme.com", "run_id": "t1", "score": 71.5, "tier": "A",
    "dimensions": {"Hiring Need": 30.0}, "explanation": ["+30 Hiring Need"],
    "scored_at": datetime(2026, 9, 16).isoformat(), "metrics": {"req_count": 5},
    "name": "Acme Data", "industry": "Software", "region": "GA", "hq_state": "GA",
    "employee_count": None, "revenue": None, "funding": None,
    "ats_provider": "greenhouse",
}


class _FakeConn:
    def close(self):
        pass


@pytest.fixture
def client():
    return TestClient(api.app)


@pytest.fixture(autouse=True)
def _stub_db(monkeypatch):
    monkeypatch.setattr(db, "connect", lambda: _FakeConn())
    monkeypatch.setattr(db, "get_scores", lambda c, r: [SCORE_ROW] if r == "t1" else [])
    monkeypatch.setattr(db, "latest_scores_per_domain", lambda c: [SCORE_ROW])
    monkeypatch.setattr(db, "contacts_by_domain", lambda c, d: {})
    monkeypatch.setattr(db, "companies_by_domain", lambda c, d: {"acme.com": Company(domain="acme.com", name="Acme Data")})
    monkeypatch.setattr(db, "get_company", lambda c, d: Company(domain="acme.com", name="Acme Data") if d == "acme.com" else None)
    monkeypatch.setattr(db, "get_postings", lambda c, d: [])
    monkeypatch.setattr(db, "get_qualification", lambda c, d, r: None)
    monkeypatch.setattr(db, "get_contacts", lambda c, d: [])


# --- Excel endpoint ----------------------------------------------------------

def test_excel_export_returns_200_with_correct_media_type(client):
    r = client.get("/runs/t1/export/excel")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_excel_export_sets_content_disposition_attachment(client):
    r = client.get("/runs/t1/export/excel")
    disposition = r.headers["content-disposition"]
    assert "attachment" in disposition
    assert "research_companies_t1.xlsx" in disposition


def test_excel_export_body_is_a_valid_workbook_with_expected_row(client):
    r = client.get("/runs/t1/export/excel")
    wb = load_workbook(BytesIO(r.content))
    ws = wb.active
    assert ws["A1"].value == "Company"
    assert ws["A2"].value == "Acme Data"


def test_excel_export_empty_run_returns_headers_only(client, monkeypatch):
    monkeypatch.setattr(db, "get_scores", lambda c, r: [])
    monkeypatch.setattr(db, "latest_scores_per_domain", lambda c: [])
    r = client.get("/runs/empty-run/export/excel")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.active.max_row == 1


# --- PDF endpoint ------------------------------------------------------------

def test_pdf_export_returns_200_with_correct_media_type(client):
    r = client.get("/runs/t1/export/pdf")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/zip"


def test_pdf_export_sets_content_disposition_attachment(client):
    r = client.get("/runs/t1/export/pdf")
    disposition = r.headers["content-disposition"]
    assert "attachment" in disposition
    assert "research_reports_t1.zip" in disposition


def test_pdf_export_body_is_a_valid_zip_with_one_pdf(client):
    r = client.get("/runs/t1/export/pdf")
    zf = zipfile.ZipFile(BytesIO(r.content))
    assert zf.namelist() == ["acme.com.pdf"]
    assert zf.read("acme.com.pdf").startswith(b"%PDF")


def test_pdf_export_empty_run_returns_valid_empty_zip(client, monkeypatch):
    monkeypatch.setattr(db, "get_scores", lambda c, r: [])
    monkeypatch.setattr(db, "latest_scores_per_domain", lambda c: [])
    r = client.get("/runs/empty-run/export/pdf")
    assert r.status_code == 200
    zf = zipfile.ZipFile(BytesIO(r.content))
    assert zf.namelist() == []
    assert zf.testzip() is None


def test_export_endpoints_never_call_the_discover_pipeline(client, monkeypatch):
    """Exports must be read-only against existing stored data -- they must
    never trigger a new research run."""
    import csg.discover as discover_mod

    def _boom(*a, **k):
        raise AssertionError("export must not call the discovery/research pipeline")

    monkeypatch.setattr(discover_mod, "discover", _boom, raising=False)
    r1 = client.get("/runs/t1/export/excel")
    r2 = client.get("/runs/t1/export/pdf")
    assert r1.status_code == 200
    assert r2.status_code == 200
